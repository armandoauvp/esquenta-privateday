#!/usr/bin/env python3
"""
Importa a lista de ingressos (XLSX ou CSV exportado da plataforma de vendas, ex.: Guru)
para o servidor da pista (Cloudflare, pasta servidor/).

O app manda o SHA-256 do e-mail digitado e o servidor confere contra esta lista. Só o
hash (com sal) sai deste computador: o e-mail em texto nunca vai pra lugar nenhum.

Pré-requisitos (uma vez):
    pip install openpyxl     (só para .xlsx)
    A chave de admin do servidor (o secret ADMIN_CHAVE, ver docs/deploy.md), passada em
    --chave-admin ou na variável de ambiente ESQUENTA_ADMIN_CHAVE.

Uso:
    python tools/importar-ingressos.py Guru-Vendas.xlsx                  # importa (soma à lista)
    python tools/importar-ingressos.py Guru-Vendas.xlsx --simular        # só mostra
    python tools/importar-ingressos.py Guru-Vendas.xlsx --substituir     # tira quem saiu (reembolso)
    python tools/importar-ingressos.py --adicionar fulano@exemplo.com    # cortesia/convidado
    python tools/importar-ingressos.py --remover fulano@exemplo.com      # tira um e-mail avulso

E-mails incluídos com --adicionar (cortesias, convidados, equipe) ficam marcados como
manuais: o --substituir nunca os apaga, mesmo que não estejam na planilha. Quem sai da
lista perde a coroa e a mesa de DJ na próxima vez que conectar.

A coluna de e-mail é detectada pelo cabeçalho (algo com "mail", ex.: "email contato");
use --coluna para indicar outra. Se houver coluna "status", só entram vendas aprovadas
("Aprovada"/"Completa"); --todos-status desliga esse filtro. No CSV, separador (vírgula
ou ponto e vírgula) e codificação são detectados. Pode rodar quantas vezes quiser.
"""

import argparse
import csv
import hashlib
import io
import json
import os
import pathlib
import re
import sys
import urllib.error
import urllib.request

# precisa bater com o index.html (INGRESSO_SAL e normalizarEmail)
SAL = "auvp-privateday:"
SERVIDOR = "https://esquenta-pista.esquenta-pista-servidor.workers.dev"
EMAIL_RE = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")


def normalizar(email: str) -> str:
    return email.strip().lower()


def hash_email(email: str) -> str:
    return hashlib.sha256((SAL + normalizar(email)).encode("utf-8")).hexdigest()


# status de venda que dão direito a ingresso (Guru: "Aprovada"; outras plataformas variam)
STATUS_OK = {"aprovada", "aprovado", "completa", "completo", "paga", "pago", "approved", "paid", "complete"}


def ler_planilha(caminho: pathlib.Path) -> tuple[list[str], list[dict]]:
    """Cabeçalho + linhas (dicts) de um .xlsx (primeira aba) ou .csv."""
    if caminho.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        aba = load_workbook(caminho, read_only=True, data_only=True).worksheets[0]
        linhas = aba.iter_rows(values_only=True)
        cabecalho = [str(c or "").strip() for c in next(linhas, ())]
        return cabecalho, [dict(zip(cabecalho, ("" if v is None else str(v) for v in l))) for l in linhas]
    bruto = caminho.read_bytes()
    for codificacao in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = bruto.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t")
    leitor = csv.DictReader(io.StringIO(texto), dialect=dialeto)
    return list(leitor.fieldnames or []), list(leitor)


def ler_emails(caminho: pathlib.Path, coluna: str | None, todos_status: bool) -> list[str]:
    cabecalho, linhas = ler_planilha(caminho)
    if coluna is None:
        candidatas = [c for c in cabecalho if c and "mail" in c.lower()]
        if not candidatas:
            sys.exit(f"ERRO: não achei coluna de e-mail no cabeçalho {cabecalho}. Use --coluna.")
        coluna = candidatas[0]
    elif coluna not in cabecalho:
        sys.exit(f"ERRO: coluna {coluna!r} não existe. Cabeçalho: {cabecalho}")
    print(f"Coluna de e-mail: {coluna!r}")
    col_status = next((c for c in cabecalho if c.lower() == "status"), None)
    if col_status and not todos_status:
        print(f"Filtrando pela coluna {col_status!r}: só vendas aprovadas/completas")

    emails, invalidos, fora_status = set(), 0, 0
    for linha in linhas:
        if col_status and not todos_status and (linha.get(col_status) or "").strip().lower() not in STATUS_OK:
            fora_status += 1
            continue
        valor = normalizar(linha.get(coluna) or "")
        if EMAIL_RE.match(valor):
            emails.add(valor)
        elif valor:
            invalidos += 1
    print(f"{len(linhas)} linha(s) lidas.")
    if fora_status:
        print(f"{fora_status} venda(s) ignorada(s) por status (cancelada, reembolsada, pendente...).")
    if invalidos:
        print(f"Aviso: {invalidos} linha(s) com e-mail inválido foram ignoradas.")
    return sorted(emails)


def enviar(servidor: str, sala: str, chave: str, corpo: dict) -> dict:
    req = urllib.request.Request(
        f"{servidor.rstrip('/')}/admin/ingressos?sala={sala}",
        data=json.dumps(corpo).encode("utf-8"),
        headers={"authorization": f"Bearer {chave}", "content-type": "application/json",
                 "user-agent": "importar-ingressos"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            return json.loads(r.read())
    except urllib.error.HTTPError as e:
        sys.exit(f"ERRO {e.code}: {e.read().decode('utf-8', 'replace')} (confira a --chave-admin)")


def main() -> None:
    # console do Windows: acentos legíveis mesmo fora de um terminal UTF-8
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description="Importa a lista de ingressos para o servidor da pista.")
    ap.add_argument("planilha", type=pathlib.Path, nargs="?", help="XLSX ou CSV exportado da plataforma de vendas")
    ap.add_argument("--adicionar", action="append", metavar="EMAIL", default=[],
                    help="inclui um e-mail avulso (cortesia/convidado); pode repetir. Não é apagado pelo --substituir")
    ap.add_argument("--remover", action="append", metavar="EMAIL", default=[],
                    help="tira um e-mail da lista (inclusive manual); pode repetir")
    ap.add_argument("--chave-admin", default=os.environ.get("ESQUENTA_ADMIN_CHAVE"),
                    help="ADMIN_CHAVE do servidor (padrão: variável ESQUENTA_ADMIN_CHAVE)")
    ap.add_argument("--servidor", default=SERVIDOR, help=f"endereço do servidor (padrão: {SERVIDOR})")
    ap.add_argument("--sala", default="principal", help="sala (padrão: principal; testes: teste)")
    ap.add_argument("--coluna", help="nome da coluna de e-mail (padrão: detecta pelo cabeçalho)")
    ap.add_argument("--todos-status", action="store_true", help="não filtra pela coluna status")
    ap.add_argument("--simular", action="store_true", help="só conta e mostra, sem enviar nada")
    ap.add_argument("--substituir", action="store_true",
                    help="tira da lista os ingressos que não estão nesta planilha (reembolsos)")
    args = ap.parse_args()
    if not args.planilha and not args.adicionar and not args.remover:
        ap.error("informe a planilha, --adicionar EMAIL ou --remover EMAIL")
    if args.substituir and not args.planilha:
        ap.error("--substituir só faz sentido com a planilha completa")

    avulsos = [normalizar(e) for e in args.adicionar]
    retirar = [normalizar(e) for e in args.remover]
    for e in avulsos + retirar:
        if not EMAIL_RE.match(e):
            sys.exit(f"ERRO: e-mail inválido: {e!r}")
    emails = ler_emails(args.planilha, args.coluna, args.todos_status) if args.planilha else []
    if args.planilha:
        print(f"{len(emails)} e-mail(s) único(s) com ingresso.")
    if avulsos:
        print(f"{len(avulsos)} e-mail(s) avulso(s) a incluir como manual.")
    if retirar:
        print(f"{len(retirar)} e-mail(s) a remover.")
    if args.simular:
        print("Simulação: nada foi enviado.")
        return
    if not args.chave_admin:
        sys.exit("ERRO: informe --chave-admin (ou a variável ESQUENTA_ADMIN_CHAVE), ou use --simular.")

    r = enviar(args.servidor, args.sala, args.chave_admin, {
        "hashes": [hash_email(e) for e in emails],
        "manuais": [hash_email(e) for e in avulsos],
        "remover": [hash_email(e) for e in retirar],
        "substituir": args.substituir,
    })
    print(f"Lista no servidor: {r['ingressos']} ingresso(s) ({r['novos']} novo(s), {r['removidos']} removido(s)).")


if __name__ == "__main__":
    main()
